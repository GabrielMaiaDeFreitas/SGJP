from io import BytesIO

from flask import Response

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment


class ExcelService:

    @staticmethod
    def exportar(
        dados,
        colunas,
        nome_arquivo,
        titulo
    ):

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = titulo

        # ==================================================
        # TÍTULO
        # ==================================================

        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(colunas)
        )

        titulo_cell = sheet["A1"]

        titulo_cell.value = titulo

        titulo_cell.font = Font(
            bold=True,
            size=16,
            color="FFFFFF"
        )

        titulo_cell.alignment = Alignment(
            horizontal="center"
        )

        titulo_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="17324F"
        )

        # ==================================================
        # CABEÇALHO
        # ==================================================

        borda = Border(

            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")

        )

        for coluna_excel, coluna in enumerate(colunas, start=1):

            cell = sheet.cell(
                row=2,
                column=coluna_excel
            )

            cell.value = coluna

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="17324F"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

            cell.border = borda

        # ==================================================
        # DADOS
        # ==================================================

        linha = 3

        for registro in dados:

            for coluna_excel, coluna in enumerate(colunas, start=1):

                cell = sheet.cell(
                    row=linha,
                    column=coluna_excel
                )

                cell.value = getattr(
                    registro,
                    coluna
                )

                cell.border = borda

            linha += 1

        # ==================================================
        # AJUSTA LARGURA
        # ==================================================

        for coluna in sheet.columns:

            tamanho = 0

            letra = coluna[0].column_letter

            for cell in coluna:

                if cell.value:

                    tamanho = max(
                        tamanho,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[letra].width = tamanho + 4

        # ==================================================
        # SALVA
        # ==================================================

        buffer = BytesIO()

        workbook.save(buffer)

        buffer.seek(0)

        return Response(

            buffer.getvalue(),

            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={

                "Content-Disposition":
                f"attachment; filename={nome_arquivo}.xlsx"

            }

        )