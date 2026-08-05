from datetime import date
from datetime import datetime
from io import BytesIO

from flask import Response

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from app.utils.object_utils import ObjectUtils

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side
)


class ExcelService:

    @staticmethod
    def formatar(valor):

        if isinstance(valor, bool):

            return "Ativo" if valor else "Inativo"

        if isinstance(valor, (date, datetime)):

            return valor.strftime("%d/%m/%Y")

        return valor

    @staticmethod
    def exportar(
        dados,
        colunas,
        nome_arquivo,
        titulo,
        labels
    ):

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = titulo

        borda = Border(

            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")

        )

        # ==================================================
        # TÍTULO
        # ==================================================

        sheet.merge_cells(

            start_row=1,

            start_column=1,

            end_row=1,

            end_column=len(colunas)

        )

        titulo_cell = sheet["A1"]

        titulo_cell.value = titulo

        titulo_cell.font = Font(

            bold=True,

            size=16,

            color="FFFFFF"

        )

        titulo_cell.fill = PatternFill(

            fill_type="solid",

            fgColor="17324F"

        )

        titulo_cell.alignment = Alignment(

            horizontal="center",

            vertical="center"

        )

        titulo_cell.border = borda

        # ==================================================
        # CABEÇALHO
        # ==================================================

        for indice, coluna in enumerate(colunas, start=1):

            cell = sheet.cell(

                row=2,

                column=indice

            )

            cell.value = labels.get(

                coluna,

                coluna

            )

            cell.font = Font(

                bold=True,

                color="FFFFFF"

            )

            cell.fill = PatternFill(

                fill_type="solid",

                fgColor="17324F"

            )

            cell.alignment = Alignment(

                horizontal="center"

            )

            cell.border = borda

        # ==================================================
        # DADOS
        # ==================================================

        linha = 3

        for registro in dados:

            for indice, coluna in enumerate(colunas, start=1):

                cell = sheet.cell(

                    row=linha,

                    column=indice

                )

                cell.value = ExcelService.formatar(

                    ObjectUtils.obter_valor(

                        registro,

                        coluna

                    )

                )

                cell.border = borda

            linha += 1

        # ==================================================
        # CONGELA O CABEÇALHO
        # ==================================================

        sheet.freeze_panes = "A3"

        # ==================================================
        # FILTRO AUTOMÁTICO
        # ==================================================

        sheet.auto_filter.ref = sheet.dimensions

        # ==================================================
        # LARGURA DAS COLUNAS
        # ==================================================

        for indice, nome_coluna in enumerate(colunas, start=1):

            letra = get_column_letter(indice)

            tamanho = len(labels.get(nome_coluna, nome_coluna))

            for linha in range(3, sheet.max_row + 1):

                valor = sheet.cell(
                    row=linha,
                    column=indice
                ).value

                if valor is not None:

                    tamanho = max(
                        tamanho,
                        len(str(valor))
                    )

            sheet.column_dimensions[letra].width = tamanho + 4

        # ==================================================
        # SALVA
        # ==================================================

        buffer = BytesIO()

        workbook.save(buffer)

        buffer.seek(0)

        return Response(

            buffer.getvalue(),

            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={

                "Content-Disposition":

                f'attachment; filename="{nome_arquivo}.xlsx"'

            }

        )